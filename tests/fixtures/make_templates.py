#!/usr/bin/env python3
"""Régénère les modèles vierges d'AuxiOffice (app/src-tauri/templates).

Styles proches des valeurs par défaut de Microsoft Office (Calibri 11,
présentations 16:9) ; LibreOffice substitue Carlito, aux métriques
identiques, quand Calibri est absent.
Dessin (AuxiDraw) : page Lettre US en portrait, comme les documents texte.
Dépendances : python-docx, openpyxl, python-pptx.
"""

import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "app/src-tauri/templates")
out.mkdir(parents=True, exist_ok=True)

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
doc.core_properties.author = ""
doc.core_properties.title = ""
doc.save(out / "vierge.docx")

wb = Workbook()
wb.active.title = "Feuille1"
wb.properties.creator = ""
wb.save(out / "vierge.xlsx")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[0])
prs.core_properties.author = ""
prs.save(out / "vierge.pptx")


def write_odg(path: Path):
    """Dessin OpenDocument minimal : une page Lettre US (21,59 × 27,94 cm),
    marges de 1 cm, langue fr-CA. LibreOffice complète le reste à l'ouverture."""
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
          'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
          'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
          'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" office:version="1.3"')
    styles = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles {ns}>
 <office:styles>
  <style:default-style style:family="graphic">
   <style:text-properties fo:language="fr" fo:country="CA" style:font-name="Liberation Sans" fo:font-size="18pt"/>
  </style:default-style>
 </office:styles>
 <office:automatic-styles>
  <style:page-layout style:name="PM1">
   <style:page-layout-properties fo:margin-top="1cm" fo:margin-bottom="1cm" fo:margin-left="1cm" fo:margin-right="1cm" fo:page-width="21.59cm" fo:page-height="27.94cm" style:print-orientation="portrait"/>
  </style:page-layout>
  <style:style style:name="Mdp1" style:family="drawing-page">
   <style:drawing-page-properties draw:background-size="border" draw:fill="none"/>
  </style:style>
 </office:automatic-styles>
 <office:master-styles>
  <style:master-page style:name="Default" style:page-layout-name="PM1" draw:style-name="Mdp1"/>
 </office:master-styles>
</office:document-styles>
"""
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content {ns}>
 <office:body>
  <office:drawing>
   <draw:page draw:name="Page 1" draw:master-page-name="Default"/>
  </office:drawing>
 </office:body>
</office:document-content>
"""
    manifest = f"""<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.graphics"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>
</manifest:manifest>
"""
    with zipfile.ZipFile(path, "w") as z:
        # « mimetype » d'abord et sans compression (exigence d'ODF).
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.graphics", compress_type=zipfile.ZIP_STORED)
        for name, data in (("content.xml", content), ("styles.xml", styles), ("META-INF/manifest.xml", manifest)):
            z.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)


write_odg(out / "vierge.odg")


def to_french(path: Path):
    """Langue des modèles : français du Canada (python-docx et python-pptx
    écrivent en-US, et LibreOffice vérifierait le texte avec le dictionnaire anglais)."""
    tmp = Path(tempfile.mkdtemp()) / path.name
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith(".xml"):
                t = data.decode("utf-8")
                t = re.sub(r'(<w:(?:lang|themeFontLang)\b[^>]*\bw:val=")en-US"', r'\1fr-CA"', t)
                t = re.sub(r'(<a:(?:rPr|defRPr|endParaRPr)\b[^>]*\blang=")en-US"', r'\1fr-CA"', t)
                data = t.encode("utf-8")
            dst.writestr(item, data)
    shutil.move(tmp, path)


def widen_layouts(path: Path, factor: float):
    """Le modèle de python-pptx est en 4:3 (10 pouces) : passé en 16:9, ses
    zones de titre et de texte restaient à gauche. Positions et largeurs des
    masques et des dispositions mises à l'échelle de la nouvelle largeur."""
    tmp = Path(tempfile.mkdtemp()) / path.name
    scale = lambda m: f'{m.group(1)}"{round(int(m.group(2)) * factor)}"'
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if re.match(r"ppt/(slideMasters|slideLayouts)/[^/]+\.xml$", item.filename):
                t = data.decode("utf-8")
                t = re.sub(r'(<a:off x=)"(\d+)"', scale, t)
                t = re.sub(r'(<a:ext cx=)"(\d+)"', scale, t)
                data = t.encode("utf-8")
            dst.writestr(item, data)
    shutil.move(tmp, path)


widen_layouts(out / "vierge.pptx", 13.333 / 10)
to_french(out / "vierge.docx")
to_french(out / "vierge.pptx")

for f in sorted(out.iterdir()):
    print(f.name, f.stat().st_size)
