#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Projet AuxiNux
"""Génère les documents de test d'AuxiNux Office (contenu connu, vérifiable).

    python3 tests/fixtures/make_fixtures.py <répertoire de sortie>

Nécessite python-docx, openpyxl, python-pptx et reportlab (voir tests/README.md).
Les fichiers ne sont pas versionnés : ils sont regénérés à chaque campagne de test."""

import sys
from pathlib import Path

MARK = "AuxiNux-Test-7Q2"   # marqueur recherché dans les conversions


def make_docx(path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.core_properties.title = "Document de test AuxiNux"
    doc.add_heading("Rapport de test AuxiWriter", level=1)
    p = doc.add_paragraph("Ce document contient le marqueur ")
    run = p.add_run(MARK)
    run.bold = True
    p.add_run(" et des caractères accentués : éèàùç ÉÈÀ œ « guillemets ».")
    doc.add_heading("Liste", level=2)
    for item in ("Premier élément", "Deuxième élément", "Troisième élément"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Tableau", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"L{r + 1}C{c + 1}"
    section = doc.sections[0]
    section.header.paragraphs[0].text = "En-tête AuxiNux"
    section.footer.paragraphs[0].text = "Pied de page AuxiNux"
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    doc.save(path)


def make_xlsx(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Poste", "Montant"])
    rows = [("Loyer", 1200), ("Électricité", 85.5), ("Internet", 60), ("Livres", 42.25)]
    for r in rows:
        ws.append(list(r))
    ws["A7"] = "Total"
    ws["B7"] = "=SUM(B2:B5)"
    ws["A8"] = "Moyenne"
    ws["B8"] = "=AVERAGE(B2:B5)"
    ws["D1"] = MARK
    second = wb.create_sheet("Notes")
    second["A1"] = "Deuxième feuille"
    wb.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Présentation AuxiPresent"
    slide.placeholders[1].text = MARK
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Deuxième diapositive"
    body = slide2.placeholders[1].text_frame
    body.text = "Point un"
    body.add_paragraph().text = "Point deux"
    slide2.notes_slide.notes_text_frame.text = "Notes du présentateur"
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    slide3.shapes.title.text = "Forme"
    box = slide3.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Zone de texte"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    prs.save(path)


def make_pdf(path, pages=3):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(1, pages + 1):
        c.setFont("Helvetica-Bold", 20)
        c.drawString(72, 760, f"Page {i} du PDF de test AuxiPDF")
        c.setFont("Helvetica", 12)
        c.drawString(72, 720, f"Marqueur : {MARK}")
        c.drawString(72, 700, "Texte sélectionnable et recherchable.")
        c.showPage()
    c.save()


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "build/test-fixtures")
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "rapport.docx")
    make_xlsx(out / "budget.xlsx")
    make_pptx(out / "presentation.pptx")
    make_pdf(out / "document.pdf")
    (out / "notes-lf.txt").write_bytes(f"Ligne 1 {MARK}\nLigne 2 éàü\nLigne 3\n".encode("utf-8"))
    (out / "notes-crlf.txt").write_bytes(f"Ligne 1 {MARK}\r\nLigne 2 éàü\r\nLigne 3\r\n".encode("utf-8"))
    (out / "journal.log").write_bytes(b"2026-10-01 12:00:00 INFO demarrage\n2026-10-01 12:00:01 WARN test\n")
    (out / "lisezmoi.md").write_text(f"# Titre\n\nTexte **gras** et `code` {MARK}.\n\n- élément\n", encoding="utf-8")
    (out / "config.json").write_text('{\n  "nom": "AuxiNux",\n  "actif": true,\n  "valeur": 42\n}\n', encoding="utf-8")
    (out / "config.yaml").write_text("nom: AuxiNux\nactif: true\nliste:\n  - un\n  - deux\n", encoding="utf-8")
    (out / "reglages.ini").write_text("[general]\nlangue = fr\n; commentaire\n", encoding="utf-8")
    (out / "donnees.xml").write_text('<?xml version="1.0"?>\n<racine><!-- note --><item id="1">AuxiNux</item></racine>\n', encoding="utf-8")
    (out / "tableau.csv").write_text("nom;valeur\nalpha;1\nbeta;2\n", encoding="utf-8")
    (out / "LISEZMOI").write_text(f"Fichier sans extension {MARK}\n", encoding="utf-8")
    (out / "latin1.txt").write_bytes("Texte Latin-1 : éàü\n".encode("latin-1"))
    print(f"fixtures générées dans {out}")


if __name__ == "__main__":
    main()
