#!/usr/bin/env python3
"""Génère les documents de test d'AuxiOffice (DOCX, XLSX, PPTX, DOCM, XLSM).

Ces documents remplacent des fichiers réels : ils couvrent les éléments
courants (styles, tableaux, images, en-têtes, formules, diapositives) et
servent aux tests du moteur et au banc de fidélité.

Dépendances (outillage de test uniquement) : python-docx, openpyxl,
python-pptx, Pillow.
Utilisation : generate.py <dossier de sortie>
"""

import re
import struct
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx.util import Pt as PPt

LOREM = (
    "AuxiOffice est la suite bureautique de l'écosystème AuxiNux. Ce paragraphe "
    "contient des accents (é, è, à, ç, ô, œ), des guillemets « français » et des "
    "nombres comme 1 234,56 $ afin de vérifier la mise en page et la césure. "
)


def make_image(path: Path) -> Path:
    img = Image.new("RGB", (480, 270), (0x1F, 0x4E, 0x79))
    d = ImageDraw.Draw(img)
    for i in range(0, 480, 40):
        d.rectangle([i, 200 - i // 4, i + 30, 270], fill=(0xF2, 0xA9, 0x00))
    d.ellipse([20, 20, 120, 120], fill=(255, 255, 255))
    img.save(path)
    return path


def docx_simple(out: Path) -> None:
    doc = Document()
    doc.add_heading("Lettre de bienvenue", level=1)
    p = doc.add_paragraph("Bonjour, ")
    p.add_run("texte en gras").bold = True
    p.add_run(", ")
    p.add_run("en italique").italic = True
    p.add_run(" et ")
    r = p.add_run("en couleur")
    r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    p.add_run(".")
    doc.add_paragraph(LOREM)
    doc.save(out / "simple.docx")


def docx_report(out: Path, img: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.header.paragraphs[0].text = "Rapport annuel AuxiNux — confidentiel"
    sec.footer.paragraphs[0].text = "Projet AuxiNux · office.auxinux.ca"
    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph("Rapport annuel 2026", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Sommaire des activités", style="Subtitle")

    for chap in range(1, 6):
        doc.add_heading(f"{chap}. Chapitre {chap}", level=1)
        for sub in range(1, 3):
            doc.add_heading(f"{chap}.{sub} Section", level=2)
            for _ in range(3):
                para = doc.add_paragraph(LOREM * 3)
                para.paragraph_format.first_line_indent = Cm(0.8)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.add_paragraph("Points clés :")
        for k in range(3):
            doc.add_paragraph(f"Élément important numéro {k + 1}", style="List Bullet")
        for k in range(3):
            doc.add_paragraph(f"Étape {k + 1}", style="List Number")
        if chap == 2:
            doc.add_picture(str(img), width=Cm(12))
            doc.add_paragraph("Figure 1 — Évolution des ventes", style="Caption")
        if chap == 3:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(["Région", "T1", "T2", "Total"]):
                table.rows[0].cells[i].text = h
            for reg in ["Québec", "Ontario", "Alberta", "Nouvelle-Écosse"] * 6:
                row = table.add_row().cells
                row[0].text = reg
                row[1].text = "12 500"
                row[2].text = "13 750"
                row[3].text = "26 250"
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    land = doc.add_section()
    land.orientation = WD_ORIENT.LANDSCAPE
    land.page_width, land.page_height = land.page_height, land.page_width
    doc.add_heading("Annexe en paysage", level=1)
    doc.add_paragraph(LOREM * 4)
    doc.save(out / "rapport.docx")


def docx_long(out: Path) -> None:
    """Environ 300 pages : mesure de l'ouverture et du défilement."""
    doc = Document()
    for chap in range(1, 61):
        doc.add_heading(f"Chapitre {chap}", level=1)
        for _ in range(14):
            doc.add_paragraph(LOREM * 4)
    doc.save(out / "long-300-pages.docx")


def xlsx_budget(out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    head = ["Poste", "Janvier", "Février", "Mars", "Total T1"]
    ws.append(head)
    thin = Side(style="thin", color="999999")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center")
    rows = [("Loyer", 1500, 1500, 1500), ("Salaires", 8200, 8350, 8400),
            ("Logiciels", 320.5, 320.5, 410.75), ("Déplacements", 640, 0, 1210.2),
            ("Formation", 0, 950, 0)]
    for i, (name, a, b, c) in enumerate(rows, start=2):
        ws.append([name, a, b, c, f"=SUM(B{i}:D{i})"])
    last = len(rows) + 1
    ws.append(["Total", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=SUM(D2:D{last})", f"=SUM(E2:E{last})"])
    for row in ws.iter_rows(min_row=2, max_row=last + 1, min_col=2, max_col=5):
        for c in row:
            c.number_format = '#,##0.00 "$"'
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[last + 1]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 18
    for col in "BCDE":
        ws.column_dimensions[col].width = 14
    ws["G2"] = "Moyenne mensuelle"
    ws["H2"] = f"=AVERAGE(B{last + 1}:D{last + 1})"
    ws["G3"] = "Plus gros poste"
    ws["H3"] = f'=INDEX(A2:A{last},MATCH(MAX(E2:E{last}),E2:E{last},0))'
    ws["G4"] = "Dépassement ?"
    ws["H4"] = f'=IF(E{last + 1}>30000,"Oui","Non")'
    ws.freeze_panes = "B2"

    chart = BarChart()
    chart.title = "Dépenses du trimestre"
    chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ws.add_chart(chart, "G7")

    data = wb.create_sheet("Données")
    data.append(["Date", "Client", "Montant", "TPS", "TVQ", "Total"])
    for i in range(2, 502):
        data.append([f"=DATE(2026,{(i % 12) + 1},{(i % 28) + 1})", f"Client {i % 37}", (i * 37) % 1000 + 0.99,
                     f"=C{i}*0.05", f"=C{i}*0.09975", f"=C{i}+D{i}+E{i}"])
    for r in data.iter_rows(min_row=2, min_col=1, max_col=1):
        r[0].number_format = "yyyy-mm-dd"
    wb.create_sheet("Vide")
    wb.save(out / "budget.xlsx")


def pptx_deck(out: Path, img: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "AuxiOffice"
    s.placeholders[1].text = "Présentation de test — fidélité PPTX"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Ordre du jour"
    body = s.placeholders[1].text_frame
    body.text = "Contexte"
    for t, lvl in [("Architecture", 0), ("Moteur LibreOfficeKit", 1), ("Interface maison", 1),
                   ("Feuille de route", 0), ("Questions", 0)]:
        p = body.add_paragraph()
        p.text, p.level = t, lvl

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Formes et image"
    colors = [(0x1F, 0x4E, 0x79), (0xF2, 0xA9, 0x00), (0xC0, 0x39, 0x2B)]
    for i, shape in enumerate([MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL, MSO_SHAPE.CHEVRON]):
        sh = s.shapes.add_shape(shape, Inches(0.8 + i * 2.6), Inches(2), Inches(2.2), Inches(1.4))
        sh.fill.solid()
        sh.fill.fore_color.rgb = PptColor(*colors[i])
        sh.text_frame.text = f"Forme {i + 1}"
        sh.text_frame.paragraphs[0].runs[0].font.size = PPt(20)
    s.shapes.add_picture(str(img), Inches(8.6), Inches(1.8), width=Inches(4))

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Graphique"
    cd = CategoryChartData()
    cd.categories = ["T1", "T2", "T3", "T4"]
    cd.add_series("2025", (12.5, 14.1, 13.9, 16.0))
    cd.add_series("2026", (13.2, 15.8, 16.4, 18.9))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(11), Inches(5.4), cd)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Tableau"
    rows, cols = 5, 4
    t = s.shapes.add_table(rows, cols, Inches(1), Inches(1.8), Inches(11), Inches(4)).table
    for c, h in enumerate(["Produit", "Prix", "Quantité", "Total"]):
        t.cell(0, c).text = h
    for r in range(1, rows):
        for c, v in enumerate([f"Article {r}", f"{r * 9.99:.2f} $", str(r * 3), f"{r * r * 29.97:.2f} $"]):
            t.cell(r, c).text = v
    prs.save(out / "presentation.pptx")


# --- Documents avec macros (DOCM, XLSM) ---------------------------------------
#
# Projet VBA minimal écrit à la main (MS-OVBA dans un fichier composé MS-CFB) :
# un module « Module1 » avec la macro AuxiBonjour. Les tests vérifient que
# l'enregistrement conserve ce projet (vbaProject.bin) et son code.

VBA_MACRO = 'Attribute VB_Name = "Module1"\r\nSub AuxiBonjour()\r\n    MsgBox "Bonjour AuxiOffice"\r\nEnd Sub\r\n'


def _ovba_compress(data: bytes) -> bytes:
    """Conteneur compressé MS-OVBA (2.4.1) sans répétitions : des littéraux
    seulement, en morceaux de 4 096 octets au plus."""
    out = bytearray(b"\x01")
    for start in range(0, len(data), 4096):
        chunk = data[start:start + 4096]
        body = bytearray()
        for i in range(0, len(chunk), 8):
            body.append(0)  # 8 littéraux
            body += chunk[i:i + 8]
        if len(body) + 2 > 4098:  # trop long compressé : morceau brut
            body = bytearray(chunk.ljust(4096, b"\0"))
            header = 0x3000 | 4095
        else:
            header = 0x8000 | 0x3000 | (len(body) + 2 - 3)
        out += struct.pack("<H", header) + body
    return bytes(out)


def _ovba_encrypt(project_id: str, data: bytes, seed: int) -> str:
    """Chiffrement des propriétés CMG, DPB et GC du flux PROJECT (2.4.3)."""
    version, key = 2, sum(project_id.encode("latin-1")) & 0xFF
    out = [seed, seed ^ version, seed ^ key]
    unenc1, enc1, enc2 = key, seed ^ key, seed ^ version
    plain = [7] * ((seed & 6) // 2) + list(struct.pack("<I", len(data))) + list(data)
    for b in plain:
        e = b ^ ((enc2 + unenc1) & 0xFF)
        out.append(e)
        enc2, enc1, unenc1 = enc1, e, b
    return "".join(f"{b:02X}" for b in out)


def _cfb(streams: dict) -> bytes:
    """Fichier composé MS-CFB version 3 ; flux de moins de 4 096 octets
    (dans le mini-flux). `streams` : chemin « A/B » → octets."""
    SECT, MINI, END, FREE, FATSECT, NOSTREAM = 512, 64, 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    # Arborescence : nœud = [nom, type, enfants, données].
    root = ["Root Entry", 5, {}, None]
    for path, data in streams.items():
        assert len(data) < 4096
        node = root
        parts = path.split("/")
        for p in parts[:-1]:
            node = node[2].setdefault(p, [p, 1, {}, None])
        node[2][parts[-1]] = [parts[-1], 2, {}, data]
    entries = []

    def add(node):
        entries.append(node)
        node.append(len(entries) - 1)  # identifiant
        for child in node[2].values():
            add(child)

    add(root)
    # Mini-flux et mini-FAT.
    mini, minifat = bytearray(), []
    for e in entries:
        if e[1] == 2:
            start = len(mini) // MINI
            n = max(1, -(-len(e[3]) // MINI))
            e.append(start)
            mini += e[3].ljust(n * MINI, b"\0")
            minifat += [start + i + 1 for i in range(n - 1)] + [END]
    # Répertoire : frères en liste chaînée triée (longueur puis majuscules).
    key = lambda n: (len(n[0]), n[0].upper())
    links = {}
    for e in entries:
        kids = sorted(e[2].values(), key=key)
        links.setdefault(e[4], {})["child"] = kids[0][4] if kids else NOSTREAM
        for a, b in zip(kids, kids[1:] + [None]):
            links.setdefault(a[4], {})["right"] = b[4] if b else NOSTREAM
    n_dir = -(-len(entries) // 4)
    n_minifat = -(-len(minifat) * 4 // SECT)
    n_mini = -(-len(mini) // SECT)
    n_fat = 1
    assert 1 + n_dir + n_minifat + n_mini <= SECT // 4
    dir_start, minifat_start, mini_start = 1, 1 + n_dir, 1 + n_dir + n_minifat
    fat = [FATSECT]
    for start, n in ((dir_start, n_dir), (minifat_start, n_minifat), (mini_start, n_mini)):
        fat += [start + i + 1 for i in range(n - 1)] + [END]
    fat += [FREE] * (SECT // 4 - len(fat))
    d = bytearray()
    for e in entries:
        name = e[0].encode("utf-16-le") + b"\0\0"
        lk = links.get(e[4], {})
        if e[1] == 5:
            start, size = (mini_start if mini else END), len(mini)
        elif e[1] == 2:
            start, size = e[5], len(e[3])
        else:
            start, size = 0, 0
        d += name.ljust(64, b"\0") + struct.pack("<HBB", len(name), e[1], 1)
        d += struct.pack("<III", NOSTREAM, lk.get("right", NOSTREAM), lk.get("child", NOSTREAM))
        d += b"\0" * 16 + struct.pack("<I", 0) + b"\0" * 16 + struct.pack("<IQ", start, size)
    while len(d) < n_dir * SECT:
        d += b"\0" * 64 + struct.pack("<HBB", 0, 0, 0) + struct.pack("<III", NOSTREAM, NOSTREAM, NOSTREAM) + b"\0" * 36 + struct.pack("<IQ", 0, 0)
    header = bytearray(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" + b"\0" * 16)
    header += struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6) + b"\0" * 6
    header += struct.pack("<IIIIIIIII", 0, n_fat, dir_start, 0, 4096, minifat_start if minifat else END, n_minifat, END, 0)
    header += struct.pack("<I", 0) + struct.pack("<I", FREE) * 108
    body = struct.pack(f"<{SECT // 4}I", *fat) + bytes(d)
    mf = struct.pack(f"<{len(minifat)}I", *minifat)
    body += mf.ljust(n_minifat * SECT, b"\xFF") + bytes(mini).ljust(n_mini * SECT, b"\0")
    return bytes(header) + body


def vba_project() -> bytes:
    """vbaProject.bin : projet « VBAProject », module standard Module1."""
    rec = lambda rid, data: struct.pack("<HI", rid, len(data)) + data
    name = b"VBAProject"
    d = rec(0x01, struct.pack("<I", 1)) + rec(0x02, struct.pack("<I", 0x0409)) + rec(0x14, struct.pack("<I", 0x0409))
    d += rec(0x03, struct.pack("<H", 1252)) + rec(0x04, name)
    d += rec(0x05, b"") + rec(0x40, b"") + rec(0x06, b"") + rec(0x3D, b"")
    d += rec(0x07, struct.pack("<I", 0)) + rec(0x08, struct.pack("<I", 0))
    d += struct.pack("<HIIH", 0x09, 4, 1, 0)  # PROJECTVERSION
    d += rec(0x0C, b"") + rec(0x3C, b"")
    d += rec(0x0F, struct.pack("<H", 1)) + rec(0x13, struct.pack("<H", 0xFFFF))
    mod = b"Module1"
    d += rec(0x19, mod) + rec(0x47, mod.decode().encode("utf-16-le"))
    d += rec(0x1A, mod) + rec(0x32, mod.decode().encode("utf-16-le"))
    d += rec(0x1C, b"") + rec(0x48, b"")
    d += rec(0x31, struct.pack("<I", 0)) + rec(0x1E, struct.pack("<I", 0)) + rec(0x2C, struct.pack("<H", 0xFFFF))
    d += struct.pack("<HI", 0x21, 0) + struct.pack("<HI", 0x2B, 0) + struct.pack("<HI", 0x10, 0)
    pid = "{9E394C0B-697E-4AEE-9FA6-446F51FB30DC}"
    # Non protégé, sans mot de passe, visible.
    cmg, dpb, gc = _ovba_encrypt(pid, struct.pack("<I", 0), 0x2D), _ovba_encrypt(pid, b"\0", 0x5F), _ovba_encrypt(pid, b"\xFF", 0x13)
    project = (
        f'ID="{pid}"\r\nModule=Module1\r\nName="VBAProject"\r\nHelpContextID="0"\r\nVersionCompatible32="393222000"\r\n'
        f'CMG="{cmg}"\r\nDPB="{dpb}"\r\nGC="{gc}"\r\n\r\n'
        "[Host Extender Info]\r\n&H00000001={3832D640-CF90-11CF-8E43-00A0C911005A};VBE;&H00000000\r\n\r\n"
        "[Workspace]\r\nModule1=0, 0, 0, 0, C\r\n"
    ).encode("latin-1")
    return _cfb({
        "VBA/dir": _ovba_compress(d),
        "VBA/_VBA_PROJECT": struct.pack("<HHBH", 0x61CC, 0xFFFF, 0, 0),
        "VBA/Module1": _ovba_compress(VBA_MACRO.encode("latin-1")),
        "PROJECT": project,
        "PROJECTwm": mod + b"\0" + mod.decode().encode("utf-16-le") + b"\0\0" + b"\0\0",
    })


def _add_vba(src: Path, dst: Path, part_dir: str, main_part: str, macro_type: str) -> None:
    """Copie un paquet OOXML en y ajoutant le projet VBA (type « macros »)."""
    rels = f"{part_dir}/_rels/{main_part.rsplit('/', 1)[-1]}.rels"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = re.sub(rf'(PartName="/{re.escape(main_part)}" ContentType=")[^"]+"', rf'\g<1>{macro_type}"', text)
                text = text.replace("<Default ", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Default ', 1)
                data = text.encode("utf-8")
            elif item.filename == rels:
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdVba1" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
                )
            zout.writestr(item, data)
        zout.writestr(f"{part_dir}/vbaProject.bin", vba_project())


def docm_macro(out: Path) -> None:
    doc = Document()
    doc.add_heading("Document avec macro", level=1)
    doc.add_paragraph("Ce document contient la macro VBA AuxiBonjour (Module1).")
    tmp = out / "macro-tmp.docx"
    doc.save(tmp)
    _add_vba(tmp, out / "macro.docm", "word", "word/document.xml", "application/vnd.ms-word.document.macroEnabled.main+xml")
    tmp.unlink()


def xlsm_macro(out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Macros"
    ws.append(["Classeur avec macro", 1, 2, "=B1+C1"])
    tmp = out / "macro-tmp.xlsx"
    wb.save(tmp)
    _add_vba(tmp, out / "macro.xlsm", "xl", "xl/workbook.xml", "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
    tmp.unlink()


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures-out")
    out.mkdir(parents=True, exist_ok=True)
    img = make_image(out / "image.png")
    docx_simple(out)
    docx_report(out, img)
    docx_long(out)
    xlsx_budget(out)
    pptx_deck(out, img)
    docm_macro(out)
    xlsm_macro(out)
    for f in sorted(out.iterdir()):
        print(f"{f.name}\t{f.stat().st_size}")


if __name__ == "__main__":
    main()
