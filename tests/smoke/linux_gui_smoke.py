#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Projet AuxiNux
"""Tests de fumée graphiques d'AuxiNux Office sous Linux (Xvfb + xdotool).

Ouvre, modifie et enregistre réellement des documents DOCX, XLSX, PPTX et texte,
ouvre un PDF (navigation, recherche), dans un profil utilisateur neuf et — si
possible — sans aucun accès réseau (namespace réseau isolé). Les captures d'écran
et un rapport JSON sont écrits dans le répertoire de sortie.

    linux_gui_smoke.py --app <arbre installé> --fixtures <dir> --out <dir>
        [--xvfb "commande Xvfb"] [--xdotool chemin] [--network-audit strace]

Variables utiles : AUXI_TEST_DISPLAY (défaut :97)."""

import argparse
import json
import os
import shutil
import signal
import subprocess
import sys
import time
import zipfile
from pathlib import Path

MARK = "AuxiNux-Test-7Q2"


class Session:
    def __init__(self, args):
        args.app = str(Path(args.app).resolve())
        args.fixtures = str(Path(args.fixtures).resolve())
        args.out = str(Path(args.out).resolve())
        self.args = args
        self.display = os.environ.get("AUXI_TEST_DISPLAY", ":97")
        self.out = Path(args.out)
        self.out.mkdir(parents=True, exist_ok=True)
        self.home = self.out / "home"
        shutil.rmtree(self.home, ignore_errors=True)
        self.home.mkdir()
        self.docs = self.out / "docs"
        shutil.rmtree(self.docs, ignore_errors=True)
        shutil.copytree(args.fixtures, self.docs)
        self.results = []
        available = subprocess.run(["locale", "-a"], capture_output=True, text=True).stdout.lower().split()
        self.xdo_locale = "en_US.UTF-8" if "en_us.utf8" in available else "C.UTF-8"
        self.xvfb = None
        self.offline = self._userns_net_available()
        self.env = dict(os.environ, DISPLAY=self.display, HOME=str(self.home), LANG="fr_FR.UTF-8",
                        XDG_CONFIG_HOME=str(self.home / ".config"), XDG_DATA_HOME=str(self.home / ".local/share"),
                        XDG_CACHE_HOME=str(self.home / ".cache"))

    # -- infrastructure -----------------------------------------------------
    @staticmethod
    def _userns_net_available():
        return subprocess.run(["unshare", "-Un", "true"], capture_output=True).returncode == 0

    def start_x(self):
        cmd = self.args.xvfb.split() + [self.display, "-screen", "0", "1600x1000x24", "-nolisten", "tcp"]
        self.xvfb = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(50):
            if self.xdo("getdisplaygeometry", check=False).returncode == 0:
                return
            time.sleep(0.2)
        raise RuntimeError("serveur X de test indisponible")

    def stop(self):
        for name in ("auxioffice-edit", "editors_helper", "auxitext"):
            subprocess.run(["pkill", "-x", name], capture_output=True)
        time.sleep(1)
        if self.xvfb:
            self.xvfb.terminate()

    def xdo(self, *args, check=True):
        # xdotool (Xlib) convertit le texte selon sa locale : il faut une locale UTF-8 connue d'Xlib
        env = dict(self.env, LANG=self.xdo_locale, LC_ALL=self.xdo_locale)
        return subprocess.run([self.args.xdotool, *map(str, args)], env=env, capture_output=True, text=True,
                              check=check)

    def launch(self, *files, audit=False):
        prefix = ["unshare", "-Un"] if self.offline else []
        if audit and self.args.network_audit:
            prefix += [self.args.network_audit, "-f", "-qq", "-s", "200", "-e", "trace=connect,sendto,sendmsg",
                       "-o", str(self.out / f"strace-{len(self.results)}.log")]
        subprocess.Popen(prefix + [str(Path(self.args.app) / "auxioffice"), *map(str, files)], env=self.env,
                         stdout=open(self.out / "app.log", "a"), stderr=subprocess.STDOUT, cwd=str(self.home))

    def wait_window(self, name, timeout=60):
        end = time.time() + timeout
        while time.time() < end:
            res = self.xdo("search", "--onlyvisible", "--name", name, check=False)
            ids = [w for w in res.stdout.split() if w]
            if ids:
                return ids[-1]
            time.sleep(0.5)
        return None

    def shot(self, name):
        from Xlib import X, display
        from PIL import Image
        d = display.Display(self.display)
        root = d.screen().root
        g = root.get_geometry()
        raw = root.get_image(0, 0, g.width, g.height, X.ZPixmap, 0xffffffff)
        image = Image.frombytes("RGB", (g.width, g.height), raw.data, "raw", "BGRX")
        image.save(self.out / f"{name}.png")
        return image

    def wait_stable(self, timeout=60, quiet=3.0):
        """Attend que l'écran ne change plus pendant « quiet » secondes (chargement terminé)."""
        from Xlib import X, display
        d = display.Display(self.display)
        root = d.screen().root
        g = root.get_geometry()
        last, since, end = None, time.time(), time.time() + timeout
        while time.time() < end:
            raw = root.get_image(0, 0, g.width, g.height, X.ZPixmap, 0xffffffff).data
            digest = hash(bytes(raw[::997]))
            if digest != last:
                last, since = digest, time.time()
            elif time.time() - since >= quiet:
                return True
            time.sleep(0.5)
        return False

    def main_window(self, timeout=60):
        """Fenêtre principale du shell (classe X « auxioffice-editors ») : la plus grande."""
        end = time.time() + timeout
        while time.time() < end:
            res = self.xdo("search", "--onlyvisible", "--classname", "auxioffice-editors", check=False)
            best, area = None, 0
            for wid in res.stdout.split():
                geo = self.xdo("getwindowgeometry", wid, check=False).stdout
                try:
                    size = [l for l in geo.splitlines() if "Geometry" in l][0].split()[1].split("x")
                    a = int(size[0]) * int(size[1])
                except (IndexError, ValueError):
                    continue
                if a > area:
                    best, area = wid, a
            if best and area > 400 * 300:
                return best
            time.sleep(0.5)
        return None

    def main_origin(self):
        wid = self.main_window(5)
        geo = self.xdo("getwindowgeometry", wid).stdout
        pos = [l for l in geo.splitlines() if "Position" in l][0].split()[1].split(",")
        return int(pos[0]), int(pos[1])

    def click(self, x, y):
        self.xdo("mousemove", x, y, "click", 1)
        time.sleep(0.6)

    def dismiss_tips(self, ox, oy, buttons):
        """Ferme les bulles « Nouveau » affichées au premier lancement de chaque éditeur
        (positions des boutons OK pour la taille de fenêtre par défaut 1200x740)."""
        for x, y in buttons:
            self.click(ox + x, oy + y)

    def record(self, name, ok, detail=""):
        self.results.append({"test": name, "ok": bool(ok), "detail": detail})
        print(f"[{'OK' if ok else 'ÉCHEC'}] {name} {detail}")

    # -- scénarios ----------------------------------------------------------
    def wait_saved(self, path, before, timeout=20):
        end = time.time() + timeout
        while time.time() < end:
            if path.stat().st_mtime > before:
                time.sleep(1.5)
                return True
            time.sleep(0.5)
        return False

    def test_startpage(self):
        self.launch()
        ok = self.main_window(90) is not None
        fonts = self.home / ".local/share/auxioffice/data/fonts/font_selection.bin"
        end = time.time() + 60
        while not fonts.exists() and time.time() < end:
            time.sleep(1)
        time.sleep(5)
        self.shot("01-accueil")
        settings = (self.home / ".local/share/auxioffice/data/settings.xml")
        ai_off = settings.exists() and "<disable-ai>1</disable-ai>" in settings.read_text()
        self.record("écran d'accueil", ok and fonts.exists(), "fenêtre « AuxiNux Office », cache des polices créé")
        self.record("IA désactivée par défaut", ai_off, str(settings))
        xdg = (self.home / ".config/auxioffice").is_dir() and (self.home / ".local/share/auxioffice").is_dir() \
            and (self.home / ".cache/auxioffice").is_dir()
        self.record("répertoires XDG", xdg, "~/.config/auxioffice, ~/.local/share/auxioffice, ~/.cache/auxioffice")

    def test_docx(self):
        doc = self.docs / "rapport.docx"
        before = doc.stat().st_mtime
        self.launch(doc)
        ok_open = self.main_window(30) is not None
        time.sleep(4)
        self.wait_stable(90)
        ox, oy = self.main_origin()
        self.dismiss_tips(ox, oy, [(79, 287), (711, 668)])
        self.click(ox + 500, oy + 363)          # titre du document
        self.xdo("key", "End")
        self.xdo("type", "--delay", "50", " MODIF-DOCX")
        self.xdo("key", "ctrl+s")
        saved = self.wait_saved(doc, before)
        self.shot("02-docx")
        content = zipfile.ZipFile(doc).read("word/document.xml").decode("utf-8")
        app = zipfile.ZipFile(doc).read("docProps/app.xml").decode("utf-8")
        self.record("DOCX ouvrir/modifier/enregistrer", ok_open and saved and "MODIF-DOCX" in content and MARK in content,
                    "texte ajouté et contenu d'origine présents")
        self.record("métadonnée Application", "AuxiNux Office" in app, app[app.find("<Application>"):app.find("</Application>") + 14])

    def test_xlsx(self):
        import openpyxl
        doc = self.docs / "budget.xlsx"
        before = doc.stat().st_mtime
        self.launch(doc)
        time.sleep(4)
        self.wait_stable(90)
        ox, oy = self.main_origin()
        self.dismiss_tips(ox, oy, [(791, 307)])
        self.click(ox + 99, oy + 405)           # cellule A10
        self.xdo("type", "--delay", "60", "Double")
        self.xdo("key", "Tab")
        self.xdo("type", "--delay", "60", "=B7*2")
        self.xdo("key", "Return")
        self.xdo("key", "ctrl+s")
        saved = self.wait_saved(doc, before)
        self.shot("03-xlsx")
        ws = openpyxl.load_workbook(doc)["Budget"]
        values = openpyxl.load_workbook(doc, data_only=True)["Budget"]
        ok = saved and ws["B10"].value == "=B7*2" and ws["B7"].value == "=SUM(B2:B5)" and abs((values["B10"].value or 0) - 2775.5) < 1e-6
        self.record("XLSX ouvrir/modifier/enregistrer + formules", ok, f"B10={ws['B10'].value} -> {values['B10'].value}")

    def test_pptx(self):
        from pptx import Presentation
        doc = self.docs / "presentation.pptx"
        before = doc.stat().st_mtime
        self.launch(doc)
        time.sleep(4)
        self.wait_stable(90)
        ox, oy = self.main_origin()
        self.dismiss_tips(ox, oy, [(79, 287), (476, 301)])
        self.click(ox + 900, oy + 377)          # titre de la diapositive 1
        self.xdo("key", "End")
        self.xdo("type", "--delay", "60", " MODIF-PPTX")
        self.xdo("key", "Escape")
        self.xdo("key", "ctrl+m")               # nouvelle diapositive
        time.sleep(1)
        self.xdo("key", "ctrl+s")
        saved = self.wait_saved(doc, before)
        self.shot("04-pptx")
        prs = Presentation(doc)
        title = prs.slides[0].shapes.title.text_frame.text if prs.slides[0].shapes.title else ""
        self.record("PPTX ouvrir/modifier/enregistrer", saved and "MODIF-PPTX" in title and len(prs.slides) == 4,
                    f"{len(prs.slides)} diapositives, titre « {title} »")

    def test_pdf(self):
        pdf = self.docs / "document.pdf"
        self.launch(pdf)
        time.sleep(4)
        self.wait_stable(90)
        # ouvert dans AuxiPDF : le shell l'inscrit dans les documents récents (type 513 = PDF)
        recents = self.home / ".local" / "share" / "auxioffice" / "recents.xml"
        text = recents.read_text(encoding="utf-8-sig", errors="replace") if recents.exists() else ""
        opened = f'path="{pdf}" type="513"' in text
        ox, oy = self.main_origin()
        self.click(ox + 600, oy + 470)
        before = self.shot("05-pdf")
        self.xdo("key", "ctrl+f")
        time.sleep(1)
        self.xdo("type", "--delay", "60", "Page 3")
        self.xdo("key", "Return")
        time.sleep(2)
        after = self.shot("05-pdf-recherche")
        # la recherche fait défiler la vue jusqu'à la page 3 : une part importante de l'écran change
        from PIL import ImageChops
        diff = ImageChops.difference(before, after).convert("L").point(lambda v: 255 if v > 24 else 0)
        changed = sum(diff.histogram()[255:]) / (diff.width * diff.height)
        self.record("PDF ouvrir/rechercher", opened and changed > 0.05,
                    f"documents récents : {'oui' if opened else 'non'}, vue modifiée par la recherche : "
                    f"{changed:.0%} de l'écran (captures 05-pdf*.png)")

    def test_text(self):
        doc = self.docs / "notes-crlf.txt"
        before = doc.stat().st_mtime
        self.launch(doc)
        wid = self.wait_window("notes-crlf\\.txt", 30)
        if wid:
            time.sleep(2)
            geo = self.xdo("getwindowgeometry", wid).stdout
            pos = [l for l in geo.splitlines() if "Position" in l][0].split()[1].split(",")
            self.click(int(pos[0]) + 300, int(pos[1]) + 200)    # zone de texte
            self.xdo("windowfocus", "--sync", wid)                # focus clavier (pas de gestionnaire de fenêtres)
            time.sleep(0.5)
            self.xdo("key", "ctrl+End")
            # Texte ASCII : xdotool perd les caractères non ASCII sur un serveur X neuf (remappage
            # clavier) ; l'UTF-8 est vérifié sur le contenu existant et par tests/auxitext.
            self.xdo("type", "--delay", "40", "Ligne ajoutee")
            self.xdo("key", "ctrl+s")
        saved = self.wait_saved(doc, before, 10)
        self.shot("06-auxitext")
        data = doc.read_bytes()
        original = (Path(self.args.fixtures) / "notes-crlf.txt").read_bytes()
        ok = bool(wid) and saved and data.startswith(original) and data.count(b"\r\n") == 3 \
            and data.count(b"\n") == 3 and data.endswith(b"Ligne ajoutee")
        self.record("AuxiText (CRLF, UTF-8)", ok, f"CRLF={data.count(b'\r\n')}")
        md = self.docs / "lisezmoi.md"
        self.launch(md)
        self.record("AuxiText Markdown", self.wait_window("lisezmoi\\.md", 30) is not None)
        self.launch(self.docs / "LISEZMOI")
        self.record("fichier sans extension -> AuxiText", self.wait_window("^LISEZMOI ", 30) is not None)

    def network_session(self):
        """Session séparée sous strace : écran d'accueil puis un document de chaque type."""
        if not self.args.network_audit:
            return
        self.stop_app()
        self.launch(audit=True)
        self.main_window(120)
        time.sleep(20)
        for name in ("rapport.docx", "budget.xlsx", "presentation.pptx", "document.pdf"):
            self.launch(self.docs / name, audit=True)
            time.sleep(25)
        time.sleep(10)
        self.stop_app()

    def test_sigterm(self):
        """Arrêt par SIGTERM (fermeture de session) : terminaison normale, sans erreur de
        segmentation ni processus du moteur restant (patch desktop-apps 0010)."""
        self.stop_app()
        prefix = ["unshare", "-Un"] if self.offline else []
        # unshare (sans --fork) puis le lanceur remplacent leur image : le processus suivi est le shell.
        proc = subprocess.Popen(prefix + [str(Path(self.args.app) / "auxioffice")], env=self.env,
                                stdout=open(self.out / "app.log", "a"), stderr=subprocess.STDOUT, cwd=str(self.home))
        started = self.main_window(120) is not None
        time.sleep(10)
        proc.send_signal(signal.SIGTERM)
        try:
            code = proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
            code = "aucune (toujours actif après 30 s)"
        time.sleep(3)
        helpers = subprocess.run(["pgrep", "-x", "editors_helper"], capture_output=True, text=True).stdout.split()
        self.record("arrêt par SIGTERM sans plantage", started and code == -signal.SIGTERM and not helpers,
                    f"fenêtre {'ouverte' if started else 'absente'}, sortie {code}, "
                    f"{len(helpers)} processus du moteur restant(s)")

    def stop_app(self):
        for name in ("auxioffice-edit", "editors_helper", "auxitext", "strace"):
            subprocess.run(["pkill", "-x", name], capture_output=True)
        time.sleep(2)

    def network_report(self):
        if not self.args.network_audit:
            return
        external = []
        for log in sorted(self.out.glob("strace-*.log")):
            for line in log.read_text(errors="replace").splitlines():
                if "connect(" in line and ("AF_INET" in line or "AF_INET6" in line):
                    if "127.0.0." not in line and "::1" not in line:
                        external.append(line.strip()[:200])
        self.record("aucune connexion réseau sortante", not external,
                    f"{len(external)} tentative(s)" + ("" if not external else " : " + external[0]))
        (self.out / "network-attempts.txt").write_text("\n".join(external) + "\n")

    def run(self):
        self.start_x()
        try:
            for t in (self.test_startpage, self.test_docx, self.test_xlsx, self.test_pptx, self.test_pdf, self.test_text,
                      self.test_sigterm):
                try:
                    t()
                except Exception as e:   # un scénario en échec n'arrête pas les autres
                    self.record(t.__name__, False, f"exception : {e}")
            self.network_session()
            self.network_report()
        finally:
            self.stop()
        report = {"offline_namespace": self.offline, "results": self.results}
        (self.out / "gui-smoke.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
        return all(r["ok"] for r in self.results)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--app", required=True)
    p.add_argument("--fixtures", required=True)
    p.add_argument("--out", required=True)
    p.add_argument("--xvfb", default="Xvfb")
    p.add_argument("--xdotool", default="xdotool")
    p.add_argument("--network-audit", default="")
    args = p.parse_args()
    sys.exit(0 if Session(args).run() else 1)


if __name__ == "__main__":
    main()
